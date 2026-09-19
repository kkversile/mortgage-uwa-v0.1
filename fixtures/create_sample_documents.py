from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).parent / "sample-documents"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_borders(cell, color="D9D9D9"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Synthetic demonstration document - not for real lending decisions")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)


def add_header(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p = doc.add_paragraph(subtitle)
    p.style = "Subtitle"
    p.runs[0].font.color.rgb = RGBColor(80, 80, 80)
    doc.add_paragraph("This document contains synthetic values for testing PDF text extraction in Mortgage UWA.")


def add_details(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    header = table.rows[0].cells
    header[0].text = "Field"
    header[1].text = "Value"
    for cell in header:
        set_cell_shading(cell, "1F4E79")
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for index, cell in enumerate(cells):
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
    doc.add_paragraph()


def save_credit_report():
    doc = Document()
    configure_document(doc)
    add_header(doc, "Synthetic Credit Report", "Credit report for PDF score extraction")
    doc.add_heading("Applicant Information", level=1)
    add_details(doc, [
        ("Applicant Name", "Alex Morgan"),
        ("Report Date", "September 11 2026"),
        ("Report Reference", "SYN-CREDIT-2026-0001"),
    ])
    doc.add_heading("Credit Summary", level=1)
    add_details(doc, [
        ("FICO Score", "720"),
        ("Credit Score", "720"),
        ("Risk Band", "Good"),
        ("Open Accounts", "4"),
        ("Credit Utilization", "30 percent"),
        ("Outstanding Debt", "$33,600.00"),
    ])
    doc.add_heading("Verification Note", level=1)
    doc.add_paragraph("This report is generated for automated document-processing tests. The score is intentionally clear and searchable in the PDF output.")
    doc.save(OUT / "sample-credit-report.docx")


def save_payslip():
    doc = Document()
    configure_document(doc)
    add_header(doc, "Synthetic Payslip", "Income document for PDF text extraction")
    doc.add_heading("Employment Information", level=1)
    add_details(doc, [
        ("Employer", "Synthetic Technologies"),
        ("Employee", "Alex Morgan"),
        ("Pay Period", "August 2026"),
        ("Employment Type", "Salaried"),
    ])
    doc.add_heading("Pay Summary", level=1)
    add_details(doc, [
        ("Gross Pay", "$8,000.00"),
        ("Federal Tax", "$1,440.00"),
        ("State Tax", "$320.00"),
        ("Net Pay", "$6,240.00"),
    ])
    doc.add_paragraph("This synthetic payslip is provided for testing only.")
    doc.save(OUT / "sample-payslip.docx")


def save_bank_statement():
    doc = Document()
    configure_document(doc)
    add_header(doc, "Synthetic Bank Statement", "Asset document for PDF text extraction")
    doc.add_heading("Account Information", level=1)
    add_details(doc, [
        ("Institution", "Synthetic Demo Bank"),
        ("Account Holder", "Alex Morgan"),
        ("Statement Period", "August 2026"),
        ("Account Reference", "SYN-BANK-2026-0001"),
    ])
    doc.add_heading("Statement Summary", level=1)
    add_details(doc, [
        ("Beginning Balance", "$82,500.00"),
        ("Deposits", "$8,000.00"),
        ("Withdrawals", "$5,500.00"),
        ("Ending Balance", "$85,000.00"),
    ])
    doc.add_paragraph("This synthetic statement is provided for testing only.")
    doc.save(OUT / "sample-bank-statement.docx")


OUT.mkdir(parents=True, exist_ok=True)
save_credit_report()
save_payslip()
save_bank_statement()
print(f"Created sample documents in {OUT}")
